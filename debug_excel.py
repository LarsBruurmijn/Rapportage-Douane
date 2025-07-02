import streamlit as st
import pandas as pd
import openpyxl
import xlrd
import io
import tempfile
import os
from pathlib import Path
import traceback
import sys

def analyze_streamlit_upload(uploaded_file):
    """Uitgebreide analyse van Streamlit uploaded file object"""
    st.subheader("🔍 Streamlit Upload Object Analysis")
    
    try:
        st.write("**Upload Object Properties:**")
        st.write(f"- Type: {type(uploaded_file)}")
        st.write(f"- Name: {uploaded_file.name}")
        st.write(f"- Size: {uploaded_file.size} bytes")
        st.write(f"- MIME type: {uploaded_file.type}")
        
        # Check if file has been read before
        current_position = uploaded_file.tell()
        st.write(f"- Current file position: {current_position}")
        
        # Get total size by seeking to end
        uploaded_file.seek(0, 2)  # Seek to end
        total_size = uploaded_file.tell()
        uploaded_file.seek(current_position)  # Reset position
        st.write(f"- Total file size: {total_size} bytes")
        
        # Read first few bytes to check if it's really Excel
        uploaded_file.seek(0)
        first_bytes = uploaded_file.read(10)
        uploaded_file.seek(0)
        st.write(f"- First 10 bytes: {first_bytes}")
        st.write(f"- First bytes as hex: {first_bytes.hex()}")
        
        # Check Excel magic numbers
        excel_magic = first_bytes[:4]
        if excel_magic == b'PK\x03\x04':
            st.success("✅ File has correct Excel/ZIP magic number")
        else:
            st.error(f"❌ File does not have Excel magic number. Found: {excel_magic}")
            
        return True
    except Exception as e:
        st.error(f"❌ Error analyzing upload object: {str(e)}")
        st.code(traceback.format_exc())
        return False

def test_multiple_read_methods(uploaded_file):
    """Test verschillende manieren om Excel bestanden te lezen"""
    st.subheader("🧪 Testing Multiple Read Methods")
    
    methods = [
        ("pandas + openpyxl", lambda f: pd.read_excel(f, engine='openpyxl')),
        ("pandas + xlrd", lambda f: pd.read_excel(f, engine='xlrd')),
        ("pandas default", lambda f: pd.read_excel(f)),
        ("direct openpyxl", lambda f: openpyxl.load_workbook(f)),
        ("openpyxl read_only", lambda f: openpyxl.load_workbook(f, read_only=True)),
        ("openpyxl data_only", lambda f: openpyxl.load_workbook(f, data_only=True))
    ]
    
    results = {}
    
    for method_name, method_func in methods:
        try:
            uploaded_file.seek(0)
            result = method_func(uploaded_file)
            
            if method_name.startswith("direct") or method_name.startswith("openpyxl"):
                # Voor openpyxl workbooks
                sheet_names = result.sheetnames
                results[method_name] = f"✅ Success - Sheets: {sheet_names}"
                result.close()
            else:
                # Voor pandas DataFrames
                results[method_name] = f"✅ Success - Shape: {result.shape}, Columns: {list(result.columns)}"
                
        except Exception as e:
            results[method_name] = f"❌ Failed: {str(e)}"
    
    for method, result in results.items():
        st.write(f"**{method}:** {result}")
    
    return results

def test_temporary_file_approach(uploaded_file):
    """Test het temporary file benadering"""
    st.subheader("💾 Testing Temporary File Approach")
    
    try:
        # Reset file pointer
        uploaded_file.seek(0)
        file_bytes = uploaded_file.read()
        st.write(f"Read {len(file_bytes)} bytes from upload")
        
        # Write to temporary file
        with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
            tmp_file.write(file_bytes)
            tmp_file_path = tmp_file.name
            
        st.write(f"Written to temporary file: {tmp_file_path}")
        st.write(f"Temp file exists: {os.path.exists(tmp_file_path)}")
        st.write(f"Temp file size: {os.path.getsize(tmp_file_path)} bytes")
        
        # Try to read from temp file
        try:
            df = pd.read_excel(tmp_file_path, engine='openpyxl')
            st.success(f"✅ Successfully read from temp file: {df.shape}")
            st.dataframe(df.head())
            return True
        except Exception as e:
            st.error(f"❌ Failed to read from temp file: {str(e)}")
            return False
        finally:
            # Cleanup
            if os.path.exists(tmp_file_path):
                os.unlink(tmp_file_path)
                
    except Exception as e:
        st.error(f"❌ Temporary file approach failed: {str(e)}")
        st.code(traceback.format_exc())
        return False

def test_bytes_io_approach(uploaded_file):
    """Test BytesIO benadering"""
    st.subheader("🧠 Testing BytesIO Approach")
    
    try:
        uploaded_file.seek(0)
        file_bytes = uploaded_file.read()
        
        # Create BytesIO object
        bytes_io = io.BytesIO(file_bytes)
        
        try:
            df = pd.read_excel(bytes_io, engine='openpyxl')
            st.success(f"✅ BytesIO approach works: {df.shape}")
            st.dataframe(df.head())
            return True
        except Exception as e:
            st.error(f"❌ BytesIO approach failed: {str(e)}")
            return False
            
    except Exception as e:
        st.error(f"❌ BytesIO setup failed: {str(e)}")
        return False

def analyze_excel_structure(uploaded_file):
    """Analyseer de interne structuur van het Excel bestand"""
    st.subheader("📊 Excel Structure Analysis")
    
    try:
        uploaded_file.seek(0)
        
        # Probeer met openpyxl om interne structuur te bekijken
        wb = openpyxl.load_workbook(uploaded_file, read_only=True)
        
        st.write("**Workbook Properties:**")
        st.write(f"- Sheet names: {wb.sheetnames}")
        st.write(f"- Number of sheets: {len(wb.sheetnames)}")
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            st.write(f"**Sheet '{sheet_name}':**")
            st.write(f"  - Max row: {ws.max_row}")
            st.write(f"  - Max column: {ws.max_column}")
            
            # Get first few cells to check content
            first_row = []
            try:
                for cell in ws[1]:
                    first_row.append(cell.value)
                st.write(f"  - First row: {first_row[:5]}...")  # Show first 5 values
            except Exception as e:
                st.write(f"  - Error reading first row: {str(e)}")
        
        wb.close()
        return True
        
    except Exception as e:
        st.error(f"❌ Excel structure analysis failed: {str(e)}")
        st.code(traceback.format_exc())
        return False

def test_file_corruption(uploaded_file):
    """Test of het bestand corrupt is"""
    st.subheader("🚨 File Corruption Test")
    
    try:
        uploaded_file.seek(0)
        file_bytes = uploaded_file.read()
        
        # Check file size
        if len(file_bytes) == 0:
            st.error("❌ File is empty!")
            return False
        
        if len(file_bytes) < 1000:
            st.warning(f"⚠️ File is very small ({len(file_bytes)} bytes) - might be corrupted")
        
        # Check ZIP structure (Excel files are ZIP files)
        import zipfile
        try:
            bytes_io = io.BytesIO(file_bytes)
            with zipfile.ZipFile(bytes_io, 'r') as zip_file:
                file_list = zip_file.namelist()
                st.success(f"✅ File has valid ZIP structure with {len(file_list)} internal files")
                st.write("Internal files:", file_list[:10])  # Show first 10
                return True
        except zipfile.BadZipFile:
            st.error("❌ File is not a valid ZIP/Excel file")
            return False
        except Exception as e:
            st.error(f"❌ ZIP structure test failed: {str(e)}")
            return False
            
    except Exception as e:
        st.error(f"❌ Corruption test failed: {str(e)}")
        return False

def comprehensive_debug():
    st.title("🔬 Comprehensive Excel Debug Tool")
    st.markdown("Upload een Excel bestand om uitgebreide diagnostiek uit te voeren")
    
    uploaded_file = st.file_uploader(
        "Upload Excel bestand voor analyse",
        type=['xlsx', 'xls'],
        key="debug_file"
    )
    
    if uploaded_file:
        st.success(f"File uploaded: {uploaded_file.name}")
        
        with st.expander("📋 System Information", expanded=False):
            st.write("**Python Version:**", sys.version)
            st.write("**Pandas Version:**", pd.__version__)
            st.write("**Openpyxl Version:**", openpyxl.__version__)
            
            try:
                st.write("**Xlrd Version:**", xlrd.__version__)
            except:
                st.write("**Xlrd Version:** Not available")
        
        # Run all tests
        tests = [
            ("Upload Object Analysis", lambda: analyze_streamlit_upload(uploaded_file)),
            ("File Corruption Check", lambda: test_file_corruption(uploaded_file)),
            ("Excel Structure Analysis", lambda: analyze_excel_structure(uploaded_file)),
            ("Multiple Read Methods", lambda: test_multiple_read_methods(uploaded_file)),
            ("Temporary File Approach", lambda: test_temporary_file_approach(uploaded_file)),
            ("BytesIO Approach", lambda: test_bytes_io_approach(uploaded_file))
        ]
        
        results = {}
        
        for test_name, test_func in tests:
            with st.expander(f"🧪 {test_name}", expanded=True):
                try:
                    result = test_func()
                    results[test_name] = "✅ Passed" if result else "❌ Failed"
                except Exception as e:
                    results[test_name] = f"💥 Error: {str(e)}"
                    st.error(f"Test crashed: {str(e)}")
                    st.code(traceback.format_exc())
        
        # Summary
        st.subheader("📊 Test Results Summary")
        for test_name, result in results.items():
            st.write(f"**{test_name}:** {result}")
            
        # Recommendations
        st.subheader("💡 Recommendations")
        
        failed_tests = [name for name, result in results.items() if "Failed" in result or "Error" in result]
        
        if not failed_tests:
            st.success("🎉 All tests passed! The file should work with standard Excel reading methods.")
        else:
            st.warning(f"⚠️ {len(failed_tests)} test(s) failed. This indicates potential issues:")
            for test in failed_tests:
                st.write(f"- {test}")
            
            st.markdown("""
            **Possible solutions:**
            1. Re-save the Excel file in Excel/LibreOffice as a new .xlsx file
            2. Remove complex formatting from the file
            3. Use the Temporary File approach if it worked
            4. Check if the file was corrupted during upload
            """)

if __name__ == "__main__":
    comprehensive_debug()