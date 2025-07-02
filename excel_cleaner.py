import pandas as pd
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import io
import tempfile
import os

def clean_excel_file(uploaded_file):
    """
    Probeert een Excel bestand te zuiveren van problematische formatting.
    Retourneert een nieuwe, schone Excel file.
    """
    try:
        # Lees data zonder formatting
        uploaded_file.seek(0)
        
        # Probeer alle mogelijke methoden om data te extraheren
        data_sheets = {}
        
        # Methode 1: Direct pandas lezen en dan opnieuw schrijven
        try:
            # Lees alle sheets
            all_sheets = pd.read_excel(uploaded_file, sheet_name=None, engine='openpyxl')
            for sheet_name, df in all_sheets.items():
                data_sheets[sheet_name] = df
        except:
            # Methode 2: Sheet voor sheet proberen
            uploaded_file.seek(0)
            wb = openpyxl.load_workbook(uploaded_file, data_only=True)
            
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                # Extraheer data als lists
                data = []
                for row in ws.iter_rows(values_only=True):
                    if any(cell is not None for cell in row):  # Skip lege rijen
                        data.append(row)
                
                if data:
                    # Maak DataFrame
                    if len(data) > 1:
                        df = pd.DataFrame(data[1:], columns=data[0])
                    else:
                        df = pd.DataFrame(data)
                    
                    data_sheets[sheet_name] = df
        
        # Schrijf naar nieuw, schoon bestand
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            for sheet_name, df in data_sheets.items():
                df.to_excel(writer, sheet_name=sheet_name, index=False)
        
        output.seek(0)
        return output, list(data_sheets.keys())
        
    except Exception as e:
        print(f"Excel cleaning mislukt: {e}")
        return None, []

def test_clean_excel():
    """Test functie"""
    print("Excel cleaner module geladen")

if __name__ == "__main__":
    test_clean_excel()