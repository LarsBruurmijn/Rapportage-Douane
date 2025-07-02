import pandas as pd
import openpyxl
import io
import time
import tracemalloc
import gc
from pathlib import Path
import numpy as np

def create_large_excel_file(rows=10000, cols=20, filename="large_test.xlsx"):
    """Creëert een groot Excel bestand voor testing"""
    print(f"🏗️  Creëert Excel bestand met {rows:,} rijen en {cols} kolommen...")
    
    # Generate data
    data = {}
    for i in range(cols):
        if i < 5:  # Text columns
            data[f'Text_Col_{i}'] = [f'Deze is een lange tekst string met veel woorden en zinnen die ruimte innemen cel {row}_{i}' for row in range(rows)]
        elif i < 10:  # Number columns  
            data[f'Number_Col_{i}'] = np.random.randint(1, 100000, rows)
        elif i < 15:  # Float columns
            data[f'Float_Col_{i}'] = np.random.random(rows) * 1000
        else:  # Mixed columns
            data[f'Mixed_Col_{i}'] = [f'Mixed_{row}_{np.random.randint(1,1000)}' for row in range(rows)]
    
    df = pd.DataFrame(data)
    
    # Save with different methods
    try:
        # Method 1: Standard pandas
        start_time = time.time()
        df.to_excel(filename, index=False, engine='openpyxl')
        pandas_time = time.time() - start_time
        
        file_size = Path(filename).stat().st_size
        print(f"✅ Bestand aangemaakt: {file_size/1024/1024:.1f} MB in {pandas_time:.1f}s")
        return filename, file_size
        
    except Exception as e:
        print(f"❌ Fout bij aanmaken: {str(e)}")
        return None, 0

def test_memory_usage_reading(filename):
    """Test memory usage bij het lezen van Excel bestanden"""
    print(f"\n🧠 Memory usage test voor: {filename}")
    
    file_size = Path(filename).stat().st_size
    print(f"📁 Bestandsgrootte: {file_size/1024/1024:.1f} MB")
    
    methods = [
        ("Pandas + Openpyxl", lambda f: pd.read_excel(f, engine='openpyxl')),
        ("Pandas + Xlrd", lambda f: try_read_xlrd(f)),
        ("Openpyxl read_only", lambda f: read_openpyxl_readonly(f)),
        ("Openpyxl data_only", lambda f: read_openpyxl_dataonly(f)),
    ]
    
    results = {}
    
    for method_name, method_func in methods:
        print(f"\n--- Testing {method_name} ---")
        
        # Start memory tracking
        tracemalloc.start()
        gc.collect()  # Clean up before test
        
        try:
            start_time = time.time()
            result = method_func(filename)
            end_time = time.time()
            
            # Get memory usage
            current, peak = tracemalloc.get_traced_memory()
            tracemalloc.stop()
            
            if hasattr(result, 'shape'):
                shape = result.shape
            else:
                shape = "N/A"
            
            results[method_name] = {
                'success': True,
                'time': end_time - start_time,
                'memory_current': current / 1024 / 1024,  # MB
                'memory_peak': peak / 1024 / 1024,  # MB
                'shape': shape,
                'memory_ratio': (peak / 1024 / 1024) / (file_size / 1024 / 1024)  # Memory/file ratio
            }
            
            print(f"✅ Success: {shape}")
            print(f"⏱️  Time: {end_time - start_time:.1f}s")
            print(f"🧠 Memory peak: {peak/1024/1024:.1f} MB")
            print(f"📊 Memory ratio: {results[method_name]['memory_ratio']:.1f}x file size")
            
        except Exception as e:
            tracemalloc.stop()
            results[method_name] = {
                'success': False,
                'error': str(e)[:100],
                'time': None,
                'memory_current': None,
                'memory_peak': None
            }
            print(f"❌ Failed: {str(e)[:100]}")
        
        gc.collect()  # Clean up after test
    
    return results

def try_read_xlrd(filename):
    """Try reading with xlrd engine"""
    try:
        return pd.read_excel(filename, engine='xlrd')
    except:
        raise Exception("xlrd not available or file not compatible")

def read_openpyxl_readonly(filename):
    """Read via openpyxl read-only mode"""
    wb = openpyxl.load_workbook(filename, read_only=True, data_only=True)
    ws = wb.active
    
    data = []
    for row in ws.iter_rows(values_only=True):
        if any(cell is not None for cell in row):
            data.append(row)
    
    wb.close()
    
    if data:
        return pd.DataFrame(data[1:], columns=data[0])
    return pd.DataFrame()

def read_openpyxl_dataonly(filename):
    """Read via openpyxl data-only mode"""
    wb = openpyxl.load_workbook(filename, data_only=True)
    ws = wb.active
    
    data = []
    for row in ws.iter_rows(values_only=True):
        if any(cell is not None for cell in row):
            data.append(row)
    
    wb.close()
    
    if data:
        return pd.DataFrame(data[1:], columns=data[0])
    return pd.DataFrame()

def test_streamlit_upload_simulation(filename):
    """Simuleer Streamlit file upload voor grote bestanden"""
    print(f"\n📤 Streamlit upload simulatie voor: {filename}")
    
    file_size = Path(filename).stat().st_size
    print(f"📁 Bestandsgrootte: {file_size/1024/1024:.1f} MB")
    
    try:
        # Read file into memory (like Streamlit upload)
        tracemalloc.start()
        
        with open(filename, 'rb') as f:
            file_bytes = f.read()
        
        print(f"📊 File geladen in memory: {len(file_bytes)/1024/1024:.1f} MB")
        
        # Test different Streamlit-like scenarios
        scenarios = [
            ("BytesIO direct", lambda: test_bytesio_read(file_bytes)),
            ("Multiple seeks", lambda: test_multiple_seeks(file_bytes)),
            ("Temporary file", lambda: test_temp_file_approach(file_bytes))
        ]
        
        for scenario_name, scenario_func in scenarios:
            try:
                start_time = time.time()
                result = scenario_func()
                end_time = time.time()
                
                current, peak = tracemalloc.get_traced_memory()
                
                print(f"✅ {scenario_name}: {result.shape if hasattr(result, 'shape') else 'OK'}")
                print(f"   Time: {end_time - start_time:.1f}s, Memory: {peak/1024/1024:.1f} MB")
                
            except Exception as e:
                print(f"❌ {scenario_name}: {str(e)[:50]}...")
        
        tracemalloc.stop()
        
    except Exception as e:
        print(f"❌ Upload simulatie gefaald: {str(e)}")

def test_bytesio_read(file_bytes):
    """Test BytesIO approach"""
    bytes_io = io.BytesIO(file_bytes)
    return pd.read_excel(bytes_io, engine='openpyxl')

def test_multiple_seeks(file_bytes):
    """Test multiple file seeks (Streamlit probleem)"""
    bytes_io = io.BytesIO(file_bytes)
    
    # First read
    df1 = pd.read_excel(bytes_io, engine='openpyxl')
    
    # Reset and read again
    bytes_io.seek(0)
    df2 = pd.read_excel(bytes_io, engine='openpyxl')
    
    return df2

def test_temp_file_approach(file_bytes):
    """Test temporary file approach"""
    import tempfile
    import os
    
    with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name
    
    try:
        return pd.read_excel(tmp_path, engine='openpyxl')
    finally:
        os.unlink(tmp_path)

def main():
    print("🔬 EXCEL BESTANDSGROOTTE EN MEMORY USAGE TEST")
    print("=" * 60)
    
    # Test verschillende groottes
    test_sizes = [
        (1000, 10, "small_1k.xlsx"),      # ~500KB
        (10000, 15, "medium_10k.xlsx"),   # ~5MB  
        (50000, 20, "large_50k.xlsx"),    # ~25MB
        # (100000, 25, "xlarge_100k.xlsx")  # ~50MB - alleen als je veel memory hebt
    ]
    
    for rows, cols, filename in test_sizes:
        print(f"\n{'='*60}")
        print(f"📊 TEST: {rows:,} rijen x {cols} kolommen")
        print(f"{'='*60}")
        
        # Create test file
        created_file, file_size = create_large_excel_file(rows, cols, filename)
        
        if created_file:
            # Test reading methods
            results = test_memory_usage_reading(created_file)
            
            # Test Streamlit simulation (alleen voor kleine bestanden)
            if file_size < 10 * 1024 * 1024:  # < 10MB
                test_streamlit_upload_simulation(created_file)
            else:
                print(f"⚠️  Streamlit simulatie overgeslagen (bestand te groot: {file_size/1024/1024:.1f} MB)")
            
            # Summary
            print(f"\n📋 SAMENVATTING voor {filename}:")
            successful_methods = [name for name, result in results.items() if result.get('success')]
            failed_methods = [name for name, result in results.items() if not result.get('success')]
            
            print(f"✅ Werkende methods: {len(successful_methods)}")
            for method in successful_methods:
                r = results[method]
                print(f"   - {method}: {r['time']:.1f}s, {r['memory_peak']:.1f}MB ({r['memory_ratio']:.1f}x)")
            
            if failed_methods:
                print(f"❌ Gefaalde methods: {len(failed_methods)}")
                for method in failed_methods:
                    print(f"   - {method}: {results[method]['error']}")
            
            # Cleanup
            try:
                Path(created_file).unlink()
                print(f"🗑️  Test bestand verwijderd: {created_file}")
            except:
                pass
        
        print(f"\n⏸️  Pauze voor memory cleanup...")
        gc.collect()
        time.sleep(1)
    
    print(f"\n🎯 CONCLUSIES:")
    print("1. Als je bestanden > 10MB hebt, kan dat memory problemen veroorzaken")
    print("2. Openpyxl read_only is meestal het meest memory-efficiënt")
    print("3. Streamlit file uploads kunnen extra memory overhead hebben")
    print("4. Memory usage kan 3-10x de bestandsgrootte zijn")

if __name__ == "__main__":
    main()